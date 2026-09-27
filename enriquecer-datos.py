#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
enriquecer-datos.py · propone los datos que le faltan a un maestro de empresas, con la fuente al lado.

Un maestro de clientes o proveedores llega con huecos: empresas sin CIF, sin
código postal, sin provincia, sin teléfono. Muchos de esos datos son públicos.
Esto los busca **sin gastar tokens** —fuentes oficiales gratuitas y la propia
web de la empresa— y saca un Excel de **propuestas**, no un fichero relleno.

PROPONE, NO RELLENA
-------------------
Cada propuesta lleva el valor, la fuente, el enlace y la confianza, y la
acepta o la rechaza el cliente en el mismo Excel. Un CIF que parece bueno y es
de otra empresa es el peor error posible de una migración: está bien formado,
pasa todas las validaciones y factura a quien no es. Por eso nada entra sin que
alguien lo mire. El Excel tiene el formato del de auditoría, y
`limpiar-datos.py` aplica las propuestas aceptadas.

SOLO EMPRESAS
-------------
Los datos de una sociedad son públicos. Buscar el domicilio o el DNI de una
persona física no lo es, y el RGPD no deja hacerlo sin motivo. Una fila con
DNI o NIE, o sin identificador y sin forma jurídica en el nombre, **no se
busca ni sale del ordenador**.

DE DÓNDE SALE CADA DATO
-----------------------
| Qué                           | Fuente                                          | Red |
|-------------------------------|-------------------------------------------------|-----|
| Provincia                     | Las dos primeras cifras del código postal       | no  |
| CP, municipio, provincia      | CartoCiudad (IGN), a partir de la dirección.    | sí  |
|                               | Solo si cuadra con lo que ya sabemos del registro|    |
| CIF, teléfono, correo         | La web de la empresa: la LSSI la obliga a       | sí  |
|                               | publicar NIF, razón social y domicilio en el    |     |
|                               | aviso legal. Respetando su robots.txt           |     |
| Web                           | El dominio de un correo que no sea genérico     | sí  |
| CIF, teléfono, CP             | Lo que la web declara de sí misma en schema.org | sí  |
|                               | (JSON-LD)                                       |     |
| ¿Son suyos ese CIF y ese CP?  | VIES, de la Comisión Europea: compara el nombre | sí  |
|                               | y el CP con los del CIF. Confirma o degrada     |     |
|                               | cada CIF propuesto; con --vies, los que ya hay  |     |
| CIF sin web ni correo propio  | La Base de Datos Nacional de Subvenciones       | sí  |
|                               | (--bdns): quien ha recibido una ayuda pública   |     |

Y sin salir del ordenador, la pestaña «Comprobaciones»: la letra del CIF
frente a la forma jurídica del nombre, y el prefijo del teléfono fijo frente a
la provincia del CP.

CÓMO SABE QUE EL DATO ES DE ELLA
--------------------------------
Lee la web por bloques. El CIF del bloque que firma la agencia («Diseño web
por…») no cuenta; el del bloque con «Tomo, Folio, Hoja» o «titular» es el del
titular. Un CIF o un teléfono que sale en las webs de dos empresas del fichero
es de la agencia, del alojamiento o del grupo, y no se propone. Y un nombre se
empareja por palabras enteras: «Inditex» no es «Blinditex».

USO
---
Dentro de un proyecto del Método de Implantación Tazuke, desde su raíz y
después de limpiar:

  python .claude/herramientas/enriquecer-datos.py
      Lee 06-Migracion/Maestros-Limpios/ y escribe
      06-Migracion/Enriquecimiento-<fecha>.xlsx.

Suelto:

  python enriquecer-datos.py proveedores.xlsx --salida propuestas/

Opciones:
  --columna "nombre=Razón social"   fuerza qué columna es cada cosa (se puede repetir).
                                    Roles: nif, nombre, direccion, cp, poblacion,
                                    provincia, telefono, email, web
  --sin-red                         solo lo que no sale del ordenador
  --vies                            comprobar en VIES si el nombre y el CP son los de cada CIF
  --bdns                            buscar en la Base de Datos Nacional de Subvenciones el CIF
                                    de las empresas que no lo tienen ni lo dan en su web. Lee
                                    antes su aviso legal: el acceso es abierto, pero pueden
                                    restringirlo ante un uso abusivo
  --incluir-dudosos                 buscar también filas sin identificador ni forma
                                    jurídica (pueden ser personas: úsalo sabiendo)
  --simular                         cuenta lo que haría, sin escribir nada
  --permitir-git                    escribir aunque la salida no esté ignorada en git

La caché de lo consultado queda en la carpeta de salida, para no repetir
consultas: lleva direcciones del cliente y no se publica.

Necesita: pandas, openpyxl.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import subprocess
import sys
import threading
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
import urllib.robotparser
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path

VERSION = "1.1.1"
AGENTE = f"tazuke-enriquecedor/{VERSION} (+https://tazuke.com)"

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:  # pragma: no cover
    pass

ENTRADA_METODO = Path("06-Migracion/Maestros-Limpios")
SALIDA_METODO = Path("06-Migracion")
CARPETA_ORIGEN = "02-datos-origen"

CARTOCIUDAD = "https://www.cartociudad.es/geocoder/api/geocoder/candidatesJsonp"
VIES = "https://ec.europa.eu/taxation_customs/vies/rest-api/ms/ES/vat/{}"
VIES_COMPARA = "https://ec.europa.eu/taxation_customs/vies/rest-api/check-vat-number"
BDNS = "https://www.infosubvenciones.es/bdnstrans/api/terceros"
PAUSA = {"cartociudad": 0.3, "web": 1.0, "vies": 1.0, "bdns": 1.0}

LETRA_DNI = "TRWAGMYFPDXBNJZSQVHLCKE"
TABLA_CIF = "JABCDEFGHI"

PROVINCIAS = {
    "01": "Álava", "02": "Albacete", "03": "Alicante", "04": "Almería", "05": "Ávila",
    "06": "Badajoz", "07": "Illes Balears", "08": "Barcelona", "09": "Burgos", "10": "Cáceres",
    "11": "Cádiz", "12": "Castellón", "13": "Ciudad Real", "14": "Córdoba", "15": "A Coruña",
    "16": "Cuenca", "17": "Girona", "18": "Granada", "19": "Guadalajara", "20": "Gipuzkoa",
    "21": "Huelva", "22": "Huesca", "23": "Jaén", "24": "León", "25": "Lleida", "26": "La Rioja",
    "27": "Lugo", "28": "Madrid", "29": "Málaga", "30": "Murcia", "31": "Navarra", "32": "Ourense",
    "33": "Asturias", "34": "Palencia", "35": "Las Palmas", "36": "Pontevedra", "37": "Salamanca",
    "38": "Santa Cruz de Tenerife", "39": "Cantabria", "40": "Segovia", "41": "Sevilla",
    "42": "Soria", "43": "Tarragona", "44": "Teruel", "45": "Toledo", "46": "Valencia",
    "47": "Valladolid", "48": "Bizkaia", "49": "Zamora", "50": "Zaragoza", "51": "Ceuta",
    "52": "Melilla",
}
# Los otros nombres con los que aparece cada provincia en un maestro.
ALIAS_PROVINCIA = {
    "01": {"araba", "alava", "araba/alava", "alava/araba", "vitoria"},
    "03": {"alacant", "alicante/alacant"}, "07": {"baleares", "islas baleares", "balears", "mallorca"},
    "12": {"castello", "castellon de la plana", "castellon/castello"},
    "15": {"la coruna", "coruna", "a coruna"}, "17": {"gerona"}, "20": {"guipuzcoa"},
    "25": {"lerida"}, "26": {"rioja", "logrono"}, "31": {"nafarroa"}, "32": {"orense"},
    "33": {"oviedo", "principado de asturias"}, "35": {"gran canaria"},
    "38": {"tenerife", "s/c de tenerife", "sc tenerife"}, "39": {"santander"},
    "46": {"valencia/valencia"}, "48": {"vizcaya"}, "30": {"region de murcia"},
}

FORMAS_JURIDICAS = {
    "sl", "sa", "slu", "sau", "sll", "slp", "sc", "scp", "sccl", "scoop", "coop", "cb", "sat",
    "aie", "ute", "sociedad", "limitada", "anonima", "unipersonal", "cooperativa", "ltd",
    "limited", "gmbh", "srl", "sarl", "inc", "llc", "bv", "sas", "slne", "fundacion", "asociacion",
    "ayuntamiento", "comunidad", "consorcio", "diputacion",
}
# Lo que sobra al comparar dos nombres de calle.
PALABRAS_VIA = {
    "calle", "c", "cl", "cll", "avda", "avenida", "av", "avd", "plaza", "pza", "pl", "plz", "paseo",
    "po", "ps", "pso", "carretera", "ctra", "cra", "camino", "cmno", "cm", "ronda", "rda",
    "travesia", "trav", "urbanizacion", "urb", "poligono", "pol", "pg", "pgno", "parque", "glorieta",
    "via", "de", "del", "la", "las", "los", "el", "y", "d", "l", "sn", "n", "no", "num", "numero",
    "bajo", "bj", "local", "loc", "piso", "puerta", "pta", "esc", "escalera", "nave", "km",
    "industrial", "ind", "edificio", "edif", "oficina", "of", "planta", "pl", "izq", "dcha",
}
CORREO_GENERICO = {
    "gmail.com", "hotmail.com", "hotmail.es", "outlook.com", "outlook.es", "yahoo.com", "yahoo.es",
    "live.com", "live.es", "icloud.com", "me.com", "msn.com", "telefonica.net", "movistar.es",
    "terra.es", "ono.com", "orange.es", "vodafone.es", "protonmail.com", "gmx.com", "gmx.es",
    "aol.com", "wanadoo.es", "yandex.com", "mail.com",
}
# Las páginas donde está el NIF: la LSSI (art. 10) obliga a publicarlo. Con
# las formas catalana, gallega y vasca, y las de las condiciones de uso.
ENLACES_LEGALES = re.compile(r"aviso[-_ ]?legal|nota[-_ ]?legal|informacion[-_ ]?legal|avis[-_ ]?legal|"
                             r"lege[-_ ]?oharra|legal|lssi|condiciones|terminos|contact|privacidad|quienes|"
                             r"empresa|nosotros|about", re.I)
# El pie que firma otra empresa: su CIF, su teléfono y su correo no son los de la web.
RE_AGENCIA = re.compile(r"dise[ñn]o\s+(y\s+desarrollo\s+)?web|desarrollad[oa]\s+por|dise[ñn]ad[oa]\s+por|"
                        r"powered\s+by|web\s+by|hecho\s+por|creado\s+por|realizad[oa]\s+por|"
                        r"alojamiento\s+web|hosting|agencia\s+(de\s+)?(marketing|web)", re.I)
# Señales de que el bloque identifica al titular de la web, que es lo que pide la LSSI.
RE_TITULAR = re.compile(r"titular|responsable|raz[oó]n\s+social|denominaci[oó]n|datos\s+identificativos|"
                        r"en\s+cumplimiento", re.I)
RE_REGISTRO = re.compile(r"registro\s+mercantil|tomo\s*[:.]?\s*\d|folio\s*[:.]?\s*\d|hoja\s*[:.]?\s*[A-Z]{0,2}-?\d",
                         re.I)

RE_CIF_TEXTO = re.compile(r"(?<![A-Z0-9])([ABCDEFGHJNPQRSUVW])[\s.\-]?(\d{2})[\s.]?(\d{3})[\s.]?(\d{2})"
                          r"[\s.\-]?([0-9A-J])(?![A-Z0-9])")
RE_EMAIL = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
RE_TEL = re.compile(r"(?<![\d+])(?:(?:\+|00)34[\s.\-]?)?[6789](?:[\s.\-]?\d){8}(?!\d)")

ROLES = ("nif", "nombre", "direccion", "cp", "poblacion", "provincia", "telefono", "email", "web")


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------

def sin_acentos(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def norm(s) -> str:
    return sin_acentos(str(s)).strip().lower()


def palabras(s: str) -> list:
    return [p for p in re.split(r"[^a-z0-9]+", norm(s)) if p]


def normalizar_id(v: str) -> str:
    return re.sub(r"[\s.\-/]", "", str(v)).upper()


def es_vacio(v) -> bool:
    return v is None or str(v).strip() == "" or norm(v) in {"nan", "null", "none", "#n/a", "?"}


def morir(msg: str) -> None:
    raise SystemExit(f"x {msg}")


def cif_valido(s: str) -> bool:
    s = normalizar_id(s)
    if not re.fullmatch(r"[ABCDEFGHJNPQRSUVW]\d{7}[0-9A-J]", s):
        return False
    d = s[1:8]
    pares = sum(int(d[i]) for i in (1, 3, 5))
    impares = sum((int(d[i]) * 2) // 10 + (int(d[i]) * 2) % 10 for i in (0, 2, 4, 6))
    control = (10 - (pares + impares) % 10) % 10
    return s[8] == str(control) or s[8] == TABLA_CIF[control]


def huella(nombre: str) -> str:
    """Nombre sin acentos, puntuación ni forma jurídica, palabras ordenadas."""
    t = re.sub(r"[^a-z0-9]+", " ", sin_acentos(str(nombre)).lower().replace("&", " y "))
    return " ".join(sorted({p for p in t.split() if p not in FORMAS_JURIDICAS and len(p) > 1}))


def nombre_aparece(nombre: str, texto: str) -> bool:
    """¿Están en el texto todas las palabras con peso del nombre, enteras?
    Palabras enteras, no trozos: «ron» no puede aparecer dentro de «Ronda»."""
    buscadas = [p for p in huella(nombre).split() if len(p) >= 3]
    if not buscadas:
        return False
    presentes = set(re.sub(r"[^a-z0-9]+", " ", sin_acentos(texto).lower()).split())
    return all(p in presentes for p in buscadas)


def parecido(a: str, b: str) -> float:
    return SequenceMatcher(None, a, b).ratio() if a and b else 0.0


def codigo_provincia(texto: str) -> str | None:
    """«Malaga», «MÁLAGA», «Vizcaya» o «Bizkaia» → su código de dos cifras."""
    n = norm(texto)
    if not n:
        return None
    for cod, nombre in PROVINCIAS.items():
        if n == norm(nombre) or n in ALIAS_PROVINCIA.get(cod, set()):
            return cod
    return None


def provincia_de_cp(cp: str) -> str | None:
    s = re.sub(r"\D", "", str(cp))
    if len(s) == 4:
        s = "0" + s
    return s[:2] if len(s) == 5 and s[:2] in PROVINCIAS else None


def mismo_sitio(a: str, b: str) -> bool:
    """¿Es la misma población? Tolera acentos, artículos y «Málaga capital»."""
    ha = " ".join(p for p in palabras(a) if p not in {"el", "la", "los", "las", "capital", "de"})
    hb = " ".join(p for p in palabras(b) if p not in {"el", "la", "los", "las", "capital", "de"})
    return bool(ha and hb) and (ha == hb or ha in hb or hb in ha or parecido(ha, hb) >= 0.85)


def nucleo_calle(direccion: str) -> str:
    return " ".join(p for p in palabras(direccion) if p not in PALABRAS_VIA and len(p) > 1 and not p.isdigit()
                    and not re.fullmatch(r"\d+[a-z]{0,2}", p))


def numero_portal(direccion: str) -> int | None:
    for m in re.finditer(r"(?<!\d)(\d{1,4})(?!\d)", str(direccion)):
        return int(m.group(1))
    return None


# ---------------------------------------------------------------------------
# Leer ficheros: todo como texto, y sin perder filas
# ---------------------------------------------------------------------------

def leer_csv(texto: str, sep: str):
    import csv
    import pandas as pd
    filas = list(csv.reader(io.StringIO(texto), delimiter=sep))
    if not filas:
        return pd.DataFrame()
    cab = [c.strip() or f"Columna {i + 1}" for i, c in enumerate(filas[0])]
    n = len(cab)
    buenas, indice = [], []
    for k, f in enumerate(filas[1:]):
        if f and len(f) <= n:
            buenas.append(f + [""] * (n - len(f)))
            indice.append(k)
    return pd.DataFrame(buenas, columns=cab, index=indice, dtype=str)


def cargar_hojas(path: Path) -> dict:
    import pandas as pd
    ext = path.suffix.lower()
    if ext in {".csv", ".tsv", ".txt"}:
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                texto = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        primera = texto.splitlines()[0] if texto else ""
        return {"(csv)": leer_csv(texto, max((";", ",", "\t", "|"), key=primera.count))}
    if ext in {".xlsx", ".xlsm", ".xls"}:
        return pd.read_excel(path, sheet_name=None, dtype=str, keep_default_na=False)
    morir(f"No sé leer ficheros {ext}: {path.name}")
    return {}


# ---------------------------------------------------------------------------
# Qué columna es cada cosa
# ---------------------------------------------------------------------------

def rol_por_nombre(nombre: str) -> str | None:
    n = norm(nombre)
    compacto = re.sub(r"[^a-z0-9]", "", n)
    p = palabras(nombre)
    if compacto in {"nif", "cif", "dni", "nie", "lictradnum", "nifcif", "cifnif", "vat", "taxid"} \
            or any(x in {"nif", "cif"} for x in p):
        return "nif"
    if "email" in compacto or any(x in {"mail", "correo"} for x in p):
        return "email"
    if any(x in {"web", "url", "website", "pagina", "www", "intrntsite"} for x in p):
        return "web"
    if any(re.fullmatch(r"(telefonos?|tlf|tfno|tel|phone|movil|moviles|cellular)\d*", x) for x in p):
        return "telefono"
    if compacto in {"cp", "codpostal", "codigopostal", "cpostal", "zipcode", "zip"} or "postal" in p:
        return "cp"
    if any(x in {"provincia", "county", "prov"} for x in p):
        return "provincia"
    if any(x in {"poblacion", "localidad", "municipio", "ciudad", "city", "pueblo"} for x in p):
        return "poblacion"
    if any(x in {"direccion", "domicilio", "calle", "address", "street", "dir"} for x in p):
        return "direccion"
    if compacto in {"nombre", "razonsocial", "cardname", "empresa", "denominacion", "nombrefiscal",
                    "cliente", "proveedor", "razon"} or "razon" in p or "social" in p:
        return "nombre"
    return None


def asignar_roles(columnas: list, forzados: dict) -> dict:
    """{rol: columna}. El primero que encaja gana; lo forzado, siempre."""
    roles = {}
    for rol, col in forzados.items():
        real = next((c for c in columnas if norm(c) == norm(col)), None)
        if real is None:
            morir(f"--columna {rol}={col}: esa columna no está en el fichero")
        roles[rol] = real
    for c in columnas:
        r = rol_por_nombre(c)
        if r and r not in roles and c not in roles.values():
            roles[r] = c
    return roles


def clasificar(nif: str, nombre: str) -> str:
    """empresa, persona o dudoso. Solo una empresa se busca fuera."""
    s = normalizar_id(nif)
    if s.startswith("ES") and len(s) == 11:
        s = s[2:]
    if re.fullmatch(r"\d{8}[A-Z]|[XYZKLM]\d{7}[A-Z]", s):
        return "persona"
    if re.fullmatch(r"[ABCDEFGHJNPQRSUVW]\d{7}[0-9A-J]", s):
        return "empresa"
    if any(p in FORMAS_JURIDICAS for p in palabras(nombre)) or re.search(r"\bs\.\s?[lac]\.?", norm(nombre)):
        return "empresa"
    return "dudoso"


# ---------------------------------------------------------------------------
# La red: con pausa, con caché y diciendo quién pregunta
# ---------------------------------------------------------------------------

class Red:
    def __init__(self, cache_path: Path | None, activa: bool = True):
        self.activa = activa
        self.cache_path = cache_path
        self.cache = {}
        self.ultima = {}
        self.consultas = Counter()
        self.robots = {}
        self._candado = threading.Lock()
        self._por_servidor = defaultdict(threading.Lock)
        self.pausa_host = {}          # Crawl-delay que pide cada web
        self.abandonados = set()      # servidores que han dicho 429 o 503
        self.antibot = set()          # servidores con protección contra robots
        if cache_path and cache_path.exists():
            try:
                self.cache = json.loads(cache_path.read_text(encoding="utf-8"))
            except Exception:
                self.cache = {}

    def guardar(self) -> None:
        if self.cache_path:
            self.cache_path.write_text(json.dumps(self.cache, ensure_ascii=False), encoding="utf-8")

    def descargar(self, url: str, tipo: str, max_bytes: int = 600_000) -> str | None:
        """Sin caché. La pausa es por servidor: la cortesía se le debe a cada
        web, no al conjunto, y una pausa global convertía 500 empresas en una
        hora de espera."""
        if not self.activa:
            return None
        host = urllib.parse.urlsplit(url).netloc
        if host in self.abandonados:
            return None
        with self._candado:
            candado_host = self._por_servidor[host]
            self.consultas[tipo] += 1
        # Varias webs a la vez, pero nunca dos consultas a la misma, y con la
        # pausa que pida su robots.txt si es mayor que la nuestra.
        with candado_host:
            pausa = max(PAUSA.get(tipo, 1.0), min(self.pausa_host.get(host, 0), 10))
            espera = pausa - (time.monotonic() - self.ultima.get(host, 0))
            if espera > 0:
                time.sleep(espera)
            try:
                req = urllib.request.Request(url, headers={"User-Agent": AGENTE, "Accept-Language": "es"})
                with urllib.request.urlopen(req, timeout=8) as r:
                    crudo = r.read(max_bytes)
                    charset = r.headers.get_content_charset() or "utf-8"
                return crudo.decode(charset, errors="replace")
            except urllib.error.HTTPError as e:
                if e.code in (429, 503):
                    # Nos pide que paremos: se para con ese servidor, no se insiste
                    self.abandonados.add(host)
                return None
            except Exception:
                return None
            finally:
                self.ultima[host] = time.monotonic()

    def post_json(self, url: str, datos: dict, tipo: str) -> dict | None:
        clave = "post:" + url + json.dumps(datos, sort_keys=True, ensure_ascii=False)
        if clave not in self.cache:
            if not self.activa:
                return None
            host = urllib.parse.urlsplit(url).netloc
            with self._candado:
                self.consultas[tipo] += 1
            with self._por_servidor[host]:
                espera = PAUSA.get(tipo, 1.0) - (time.monotonic() - self.ultima.get(host, 0))
                if espera > 0:
                    time.sleep(espera)
                try:
                    req = urllib.request.Request(url, data=json.dumps(datos).encode("utf-8"), method="POST",
                                                 headers={"User-Agent": AGENTE, "Content-Type": "application/json"})
                    with urllib.request.urlopen(req, timeout=15) as r:
                        self.cache[clave] = json.loads(r.read().decode("utf-8"))
                except Exception:
                    return None          # un fallo de red no se guarda: se reintenta otro día
                finally:
                    self.ultima[host] = time.monotonic()
        return self.cache[clave]

    def get(self, url: str, tipo: str, max_bytes: int = 600_000) -> str | None:
        if url not in self.cache:
            if not self.activa:
                return None
            self.cache[url] = self.descargar(url, tipo, max_bytes)
        return self.cache[url]

    def pagina(self, url: str) -> dict | None:
        """Una página web ya reducida a su texto y sus enlaces: es lo que se
        guarda en la caché, no el HTML entero."""
        clave = "pagina:" + url
        if clave not in self.cache:
            if not self.activa:
                return None
            html = self.descargar(url, "web")
            if html is not None and RE_ANTIBOT.search(html[:20_000]) and len(texto_de_html(html)) < 3000:
                # Una protección contra robots: no se intenta esquivar. Se deja.
                self.antibot.add(urllib.parse.urlsplit(url).netloc)
                html = None
            self.cache[clave] = None if html is None else {
                "texto": texto_de_html(html)[:150_000],
                "jsonld": organizaciones_jsonld(html)[:5],
                # Solo los enlaces que pueden llevar a aviso legal o contacto, pero
                # todos: el del aviso legal suele ser de los últimos de la página.
                "enlaces": sorted({h for h in re.findall(r"""href=["']([^"'#]+)["']""", html, re.I)
                                   if ENLACES_LEGALES.search(h) and not RECURSO_ESTATICO.search(h)},
                                  key=prioridad_enlace)[:50]}
        return self.cache[clave]

    def permitido(self, url: str) -> bool:
        """Lo que el robots.txt de la web deje leer, y nada más."""
        partes = urllib.parse.urlsplit(url)
        base = f"{partes.scheme}://{partes.netloc}"
        if base not in self.robots:
            rp = urllib.robotparser.RobotFileParser()
            texto = self.get(base + "/robots.txt", "web", 100_000)
            rp.parse((texto or "").splitlines())
            self.robots[base] = rp
            demora = rp.crawl_delay(AGENTE) or rp.crawl_delay("*")
            if demora:
                self.pausa_host[partes.netloc] = float(demora)
        return self.robots[base].can_fetch(AGENTE, url)


# ---------------------------------------------------------------------------
# Las fuentes
# ---------------------------------------------------------------------------

def cartociudad(red: Red, direccion: str, cp: str, poblacion: str, provincia: str) -> tuple:
    """(candidato, confianza) o (None, motivo). Un candidato solo vale si cuadra
    con lo que ya sabemos: CartoCiudad es tolerante con las erratas, y a veces
    tanto que devuelve una calle parecida en otra provincia."""
    nucleo = nucleo_calle(direccion)
    if len(nucleo) < 3:
        return None, "la dirección no tiene nombre de calle reconocible"
    lugar = poblacion or provincia or ""
    q = ", ".join(x for x in (direccion, cp, lugar) if x)
    texto = red.get(f"{CARTOCIUDAD}?{urllib.parse.urlencode({'q': q, 'limit': 5})}", "cartociudad")
    if not texto:
        return None, "sin respuesta de CartoCiudad" if red.activa else "sin red"
    try:
        candidatos = json.loads(texto[texto.find("(") + 1:texto.rfind(")")])
    except ValueError:
        return None, "respuesta de CartoCiudad ilegible"

    prov_cp = provincia_de_cp(cp) if cp else None
    prov_txt = codigo_provincia(provincia) if provincia else None
    portal = numero_portal(direccion)
    for c in candidatos or []:
        if c.get("type") not in {"portal", "callejero"}:
            continue
        calle = re.sub(r",.*$", "", c.get("address") or "")
        calle = re.sub(rf"^{re.escape(c.get('tip_via') or '')}\s+", "", calle, flags=re.I)
        if parecido(nucleo, nucleo_calle(calle)) < 0.8 and nucleo not in nucleo_calle(calle):
            continue
        cod = str(c.get("provinceCode") or "")
        if (prov_cp and cod != prov_cp) or (prov_txt and cod != prov_txt):
            continue
        if poblacion and not (mismo_sitio(poblacion, c.get("muni") or "")
                              or mismo_sitio(poblacion, c.get("poblacion") or "")):
            continue
        cps = str(c.get("postalCode") or "").split()
        if cp and cps and re.sub(r"\D", "", cp).zfill(5) not in cps:
            continue
        contrastes = sum(bool(x) for x in (prov_cp or prov_txt, poblacion, cp))
        exacto = c.get("type") == "portal" and portal is not None and c.get("portalNumber") == portal
        if contrastes >= 1 and exacto:
            return c, "alta"
        return c, "media" if contrastes >= 1 else "baja"
    return None, "ningún candidato cuadra con el CP, la población o la provincia del registro"


def vies(red: Red, cif: str) -> bool | None:
    texto = red.get(VIES.format(normalizar_id(cif)), "vies")
    if not texto:
        return None
    try:
        return bool(json.loads(texto).get("isValid"))
    except ValueError:
        return None


def consultas_bdns(nombre: str) -> list:
    """Qué preguntarle a la BDNS. Busca la frase por su principio y no casa la
    Ñ con la N: «TELEFONICA DE ESPA» encuentra y «TELEFONICA DE ESPAÑA» no. Se
    pregunta por el nombre hasta su segunda palabra con peso, cortado antes de
    la primera Ñ, y si no sale nada, por la primera palabra sola."""
    vacias = {"DE", "DEL", "LA", "LAS", "LOS", "EL", "Y", "I", "E"}
    trozos, con_peso = [], 0
    for p in re.sub(r"[^A-ZÑ0-9 ]+", " ", sin_acentos(nombre.upper().replace("Ñ", "\0")).replace("\0", "Ñ")).split():
        if p in LETRA_DE_FORMA or len(p) == 1:      # «S» de «S. A.»: empieza la forma jurídica
            break
        if "Ñ" in p:
            if p.index("Ñ") >= 3:
                trozos.append(p[:p.index("Ñ")])
            break
        trozos.append(p)
        con_peso += p not in vacias
        if con_peso == 2:
            break
    primera = next((p for p in trozos if p not in vacias), "")
    salida = []
    for q in (" ".join(trozos).strip(), primera):
        if len(q) >= 4 and q not in salida:
            salida.append(q)
    return salida


def bdns_candidatos(red: Red, nombre: str) -> tuple:
    """(cif, nombre en BDNS, similitud) del mejor candidato, o (None, motivo, 0).

    La Base de Datos Nacional de Subvenciones publica quién ha recibido una
    ayuda, con su NIF. Con el Kit Digital y las ayudas COVID son muchas pymes.
    Es la única fuente oficial y gratuita que da un CIF a partir de un nombre.
    Solo se aceptan personas jurídicas, y solo si un único CIF se parece de
    verdad al nombre: con dos candidatos parecidos, no se propone ninguno."""
    consultas = consultas_bdns(nombre)
    if not consultas:
        return None, "nombre demasiado corto para buscarlo", 0
    texto = ""
    for q in consultas:
        texto = red.get(f"{BDNS}?{urllib.parse.urlencode({'vpd': 'GE', 'busqueda': q})}", "bdns")
        if texto is None:
            return None, "sin respuesta de la BDNS" if red.activa else "sin red", 0
        if texto.strip():
            break
    if not texto.strip():                # 204: no hay nadie con ese nombre
        return None, "ningún beneficiario de la BDNS con ese nombre", 0
    try:
        terceros = json.loads(texto).get("terceros", [])
    except ValueError:
        return None, "respuesta de la BDNS ilegible", 0
    h = huella(nombre)
    buscadas = [p for p in h.split() if len(p) >= 3]
    forma = forma_del_nombre(nombre)
    mejor = {}
    for t in terceros:
        cif, _, nombre_bdns = str(t.get("descripcion", "")).partition(" - ")
        cif = normalizar_id(cif)
        if not cif_valido(cif):          # un DNI o un NIE no se toca: persona física
            continue
        # Todas las palabras del nombre, enteras: «Inditex» no es «Blinditex»
        presentes = set(huella(nombre_bdns).split())
        if not buscadas or not all(p in presentes for p in buscadas):
            continue
        # Y la letra del CIF tiene que ser la de su forma jurídica: una SA no es una SL
        if forma and cif[0] != LETRA_DE_FORMA[forma]:
            continue
        sim = parecido(h, huella(nombre_bdns))
        if sim > mejor.get(cif, ("", 0))[1]:
            mejor[cif] = (nombre_bdns.strip(), sim)
    buenos = sorted(((c, n, s) for c, (n, s) in mejor.items() if s >= 0.85), key=lambda x: -x[2])
    if not buenos:
        return None, "ningún beneficiario de la BDNS con un nombre parecido", 0
    if len(buenos) > 1 and buenos[1][2] >= buenos[0][2] - 0.05:
        return None, f"varios CIF en la BDNS con nombres parecidos ({buenos[0][0]}, {buenos[1][0]})", 0
    return buenos[0]


FORMAS_LARGAS = [("SOCIEDAD LIMITADA LABORAL", "SLL"), ("SOCIEDAD LIMITADA NUEVA EMPRESA", "SLNE"),
                 ("SOCIEDAD LIMITADA PROFESIONAL", "SLP"), ("SOCIEDAD LIMITADA UNIPERSONAL", "SLU"),
                 ("SOCIEDAD ANONIMA UNIPERSONAL", "SAU"), ("SOCIEDAD LIMITADA", "SL"),
                 ("SOCIEDAD ANONIMA", "SA"), ("SOCIEDAD COOPERATIVA", "SCOOP"), ("COMUNIDAD DE BIENES", "CB")]


def nombre_corto(nombre: str) -> str:
    """«Talleres Pérez, Sociedad Limitada» → «TALLERES PEREZ SL». VIES compara
    el nombre de forma aproximada, pero con la forma jurídica larga dice que no
    coincide: hay que mandarla abreviada."""
    s = re.sub(r"[^A-Z0-9 ]+", " ", sin_acentos(nombre).upper())
    s = re.sub(r"\s+", " ", s).strip()
    for larga, corta in FORMAS_LARGAS:
        s = re.sub(rf"\b{larga}\b", corta, s)
    # «S L» y «S A U», de haber escrito «S. L.». Solo las formas conocidas: juntar
    # a ciegas las letras sueltas convertía «S.M.E., S.A.» (Renfe) en «S MESA».
    for suelta, junta in (("S L N E", "SLNE"), ("S L L", "SLL"), ("S L P", "SLP"), ("S L U", "SLU"),
                          ("S A U", "SAU"), ("S M E", "SME"), ("S L", "SL"), ("S A", "SA"), ("C B", "CB")):
        s = re.sub(rf"\b{suelta}\b", junta, s)
    return s


def vies_compara(red: Red, cif: str, nombre: str, cp: str) -> dict | None:
    """Pregunta a VIES si ese nombre y ese CP son los del CIF. Para España VIES
    no devuelve nombre ni dirección, pero **sí compara** lo que se le manda, y
    tolera erratas. Solo conoce a los operadores intracomunitarios: que un CIF
    no esté no significa que sea falso."""
    datos = {"countryCode": "ES", "vatNumber": normalizar_id(cif).removeprefix("ES"),
             "traderName": nombre_corto(nombre)}
    cp = re.sub(r"\D", "", cp or "")
    if len(cp) == 5:
        datos["traderPostalCode"] = cp
    r = red.post_json(VIES_COMPARA, datos, "vies")
    if not r or "valid" not in r:
        return None
    return {"valido": bool(r.get("valid")), "nombre": r.get("traderNameMatch"),
            "cp": r.get("traderPostalCodeMatch") if "traderPostalCode" in datos else None}


# La forma jurídica que dice el nombre, y la letra que le toca al CIF.
LETRA_DE_FORMA = {"SA": "A", "SAU": "A", "SL": "B", "SLU": "B", "SLL": "B", "SLNE": "B", "SLP": "B",
                  "SCOOP": "F", "COOP": "F", "CB": "E"}


def forma_del_nombre(nombre: str) -> str | None:
    ultimas = nombre_corto(nombre).split()[-2:]
    for p in reversed(ultimas):
        if p in LETRA_DE_FORMA:
            return p
    return None


# Prefijos de los teléfonos fijos por provincia (plan nacional de numeración).
# Ceuta y Melilla comparten prefijo con Cádiz y Málaga.
PREFIJO_PROVINCIA = {
    "91": {"28"}, "93": {"08"},
    "920": {"05"}, "921": {"40"}, "922": {"38"}, "923": {"37"}, "924": {"06"}, "925": {"45"},
    "926": {"13"}, "927": {"10"}, "928": {"35"}, "941": {"26"}, "942": {"39"}, "943": {"20"},
    "944": {"48"}, "945": {"01"}, "946": {"48"}, "947": {"09"}, "948": {"31"}, "949": {"19"},
    "950": {"04"}, "951": {"29"}, "952": {"29", "52"}, "953": {"23"}, "954": {"41"}, "955": {"41"},
    "956": {"11", "51"}, "957": {"14"}, "958": {"18"}, "959": {"21"}, "960": {"46"}, "961": {"46"},
    "962": {"46"}, "963": {"46"}, "964": {"12"}, "965": {"03"}, "966": {"03"}, "967": {"02"},
    "968": {"30"}, "969": {"16"}, "971": {"07"}, "972": {"17"}, "973": {"25"}, "974": {"22"},
    "975": {"42"}, "976": {"50"}, "977": {"43"}, "978": {"44"}, "979": {"34"}, "980": {"49"},
    "981": {"15"}, "982": {"27"}, "983": {"47"}, "984": {"33"}, "985": {"33"}, "986": {"36"},
    "987": {"24"}, "988": {"32"},
}


def provincias_de_telefono(tel: str) -> set | None:
    """Las provincias posibles de un fijo, o None si es móvil, 90x o no se sabe."""
    s = re.sub(r"\D", "", str(tel))[-9:]
    if len(s) != 9 or not s.startswith("9") or s.startswith("90"):
        return None
    return PREFIJO_PROVINCIA.get(s[:3]) or PREFIJO_PROVINCIA.get(s[:2])


RECURSO_ESTATICO = re.compile(r"\.(css|js|json|xml|png|jpe?g|gif|svg|webp|ico|woff2?|ttf|zip|mp4)$|/wp-(content|includes|json)/",
                              re.I)


def prioridad_enlace(url: str) -> tuple:
    """El aviso legal primero (lleva el NIF por ley), después contacto."""
    ruta = urllib.parse.urlsplit(url).path.lower()
    for i, patron in enumerate((r"aviso|legal", r"contact", r"privacidad|privacy", r"quienes|empresa|nosotros|about")):
        if re.search(patron, ruta):
            return (i, len(ruta))
    return (9, len(ruta))


RE_BLOQUE = re.compile(r"(?i)<br\s*/?>|</?(p|div|li|tr|h[1-6]|address|footer|header|section|article|"
                       r"dd|dt|ul|ol|table|nav|aside|blockquote)\b[^>]*>")
RE_ANTIBOT = re.compile(r"just a moment|cf-chl|challenge-platform|captcha|attention required|"
                        r"verify you are human|comprobando (su|tu) navegador", re.I)


def texto_de_html(html: str) -> str:
    """El texto de la página, **una línea por bloque**: así se sabe qué va
    con qué. Un CIF en la misma línea que «Tomo, Folio, Hoja» es del titular;
    uno en la línea de «Diseño web por» es de la agencia. La idea de trabajar
    por bloques es la de los extractores de Impressum alemanes."""
    import html as html_lib
    html = re.sub(r"(?is)<(script|style|noscript|svg|template|head)[^>]*>.*?</\1>", " ", html)
    html = RE_BLOQUE.sub("\n", html)
    html = html_lib.unescape(re.sub(r"(?s)<[^>]+>", " ", html))
    lineas = (re.sub(r"[^\S\n]+", " ", linea).strip() for linea in html.split("\n"))
    return "\n".join(linea for linea in lineas if linea)


def organizaciones_jsonld(html: str) -> list:
    """Los datos schema.org que la propia web declara de sí misma en JSON-LD
    (Organization, LocalBusiness…): teléfono, correo, CIF y dirección sin
    adivinar nada. Lo que hace extruct, con la librería estándar."""
    salida = []
    for bloque in re.findall(r"""(?is)<script[^>]+application/ld\+json[^>]*>(.*?)</script>""", html):
        try:
            datos = json.loads(bloque.strip())
        except ValueError:
            continue
        pendientes = datos if isinstance(datos, list) else [datos]
        while pendientes:
            d = pendientes.pop()
            if not isinstance(d, dict):
                continue
            pendientes.extend(x for x in d.get("@graph", []) if isinstance(x, dict))
            tipo = d.get("@type")
            tipos = tipo if isinstance(tipo, list) else [tipo]
            if not any(str(t).endswith(("Organization", "Business", "Corporation", "Store", "Service"))
                       for t in tipos):
                continue
            dire = d.get("address") if isinstance(d.get("address"), dict) else {}
            salida.append({k: str(v).strip() for k, v in {
                "nombre": d.get("legalName") or d.get("name"), "cif": d.get("vatID") or d.get("taxID"),
                "telefono": d.get("telephone"), "email": d.get("email"),
                "cp": dire.get("postalCode"), "poblacion": dire.get("addressLocality"),
                "provincia": dire.get("addressRegion")}.items() if v and isinstance(v, (str, int))})
    return salida


def leer_web(red: Red, web: str) -> dict:
    """La portada y hasta tres páginas legales o de contacto de la misma web.
    Devuelve {cifs, telefonos, emails, paginas, texto}."""
    if not re.match(r"https?://", web, re.I):
        web = "https://" + web.strip().lstrip("/")
    partes = urllib.parse.urlsplit(web)
    if not partes.netloc:
        return {}
    base = f"{partes.scheme}://{partes.netloc}"
    visitadas, pendientes, textos, jsonld = [], [web], [], []
    while pendientes and len(visitadas) < 4:
        url = pendientes.pop(0)
        if url in visitadas or not red.permitido(url):
            visitadas.append(url)
            continue
        pag = red.pagina(url)
        visitadas.append(url)
        if not pag:
            continue
        textos.append((url, pag["texto"]))
        jsonld += [dict(o, url=url) for o in pag.get("jsonld", [])]
        if len(visitadas) == 1:
            # Todos los enlaces de la portada, y después los mejores: el aviso
            # legal suele estar al pie, detrás de todo lo demás.
            candidatos = set()
            for href in pag["enlaces"]:
                destino = urllib.parse.urljoin(url, href).split("?")[0]
                ruta = urllib.parse.urlsplit(destino).path
                if urllib.parse.urlsplit(destino).netloc == partes.netloc and ENLACES_LEGALES.search(ruta) \
                        and not RECURSO_ESTATICO.search(ruta) and destino != url:
                    candidatos.add(destino)
            pendientes = sorted(candidatos, key=prioridad_enlace)[:3]
    cifs, tels, correos = Counter(), Counter(), Counter()
    fuente, contexto, senales = {}, {}, {}
    for url, t in textos:
        lineas = t.split("\n")
        for i, linea in enumerate(lineas):
            cerca = "\n".join(lineas[max(0, i - 3):i + 4])      # el bloque y sus vecinos
            junto = "\n".join(lineas[max(0, i - 1):i + 2])      # la línea y las dos de al lado
            agencia = bool(RE_AGENCIA.search(junto))
            for m in RE_CIF_TEXTO.finditer(linea.upper()):
                c = "".join(m.groups())
                if not cif_valido(c):
                    continue
                s = senales.setdefault(c, {"agencia": False, "titular": False, "registro": False, "puntos": 0})
                if agencia:
                    s["agencia"] = True          # el de quien firma el pie, no el de la empresa
                    continue
                cifs[c] += 1
                fuente.setdefault(("cif", c), url)
                contexto.setdefault(c, cerca)
                s["titular"] |= bool(RE_TITULAR.search(cerca))
                s["registro"] |= bool(RE_REGISTRO.search(cerca))
                # En la misma línea pesa el doble que en la de al lado: dos CIF
                # en líneas seguidas no pueden llevarse la misma señal.
                puntos = 2 * bool(RE_TITULAR.search(linea)) + 2 * bool(RE_REGISTRO.search(linea))                     + bool(RE_TITULAR.search(junto)) + bool(RE_REGISTRO.search(junto))
                s["puntos"] = max(s["puntos"], puntos)
            if agencia:
                continue
            for m in RE_TEL.finditer(linea):
                tel = re.sub(r"\D", "", m.group(0))[-9:]
                if len(tel) == 9:
                    tels[tel] += 1
                    fuente.setdefault(("tel", tel), url)
            for m in RE_EMAIL.finditer(linea):
                e = m.group(0).lower().rstrip(".")
                if not re.search(r"\.(png|jpe?g|gif|webp|svg)$", e):
                    correos[e] += 1
                    fuente.setdefault(("email", e), url)
    return {"cifs": cifs, "telefonos": tels, "emails": correos, "fuente": fuente, "base": base,
            "contexto": contexto, "senales": senales, "jsonld": jsonld,
            "texto": "\n".join(t for _, t in textos)}


# ---------------------------------------------------------------------------
# Enriquecer
# ---------------------------------------------------------------------------

@dataclass
class Propuesta:
    fila: int
    identificador: str
    columna: str
    valor_actual: str
    propuesto: str
    confianza: str
    fuente: str
    porque: str


@dataclass
class Resultado:
    entidad: str
    fichero: str
    filas: int = 0
    clases: Counter = field(default_factory=Counter)
    propuestas: list = field(default_factory=list)
    sin_columna: Counter = field(default_factory=Counter)
    vies: list = field(default_factory=list)
    notas: list = field(default_factory=list)
    comprobaciones: list = field(default_factory=list)   # (fila, identificador, qué, resultado, fuente)


def web_a_leer(fila, roles: dict, op) -> str:
    """La web que habrá que leer para esta fila, o nada. La misma decisión que
    toma `enriquecer`, adelantada para leerlas todas en paralelo."""
    def v(rol):
        c = roles.get(rol)
        return "" if c is None or es_vacio(fila[c]) else str(fila[c]).strip()
    clase = clasificar(v("nif"), v("nombre"))
    if clase == "persona" or (clase == "dudoso" and not op.incluir_dudosos) or not (v("nif") or v("nombre")):
        return ""
    web, email = v("web"), v("email")
    dominio = email.split("@")[-1].lower().strip(" ;,") if "@" in email else ""
    deducida = not web and dominio and dominio not in CORREO_GENERICO and "." in dominio
    if deducida:
        web = "https://" + dominio
    if web and (not v("nif") or not v("telefono") or not email or deducida):
        return web
    return ""


def enriquecer(df, entidad: str, fichero: str, roles: dict, red: Red, op) -> Resultado:
    res = Resultado(entidad, fichero, len(df))
    cols = list(df.columns)
    fila_origen = next((c for c in cols if norm(c) == "_fila_origen"), None)
    res.notas.append("Columnas: " + ", ".join(f"{r}={c}" for r, c in roles.items()) if roles
                     else "No se ha reconocido ninguna columna")
    if "nombre" not in roles and "nif" not in roles:
        res.notas.append("Sin columna de nombre ni de identificador: no se puede saber qué fila es una empresa")
        return res
    # Un listado de facturas también tiene nombre de cliente, pero repetido en
    # cada documento: no es un maestro, y enriquecerlo es consultar cien veces
    # lo mismo. Se reconoce porque los nombres se repiten.
    clave = roles.get("nif") or roles["nombre"]
    llenos = [str(v).strip() for v in df[clave] if not es_vacio(v)]
    if len(llenos) > 20 and len(set(llenos)) / len(llenos) < 0.5:
        res.notas.append(f"No parece un maestro: {len(set(llenos))} valores distintos de «{clave}» en "
                         f"{len(llenos)} filas (una fila por documento). No se enriquece; se enriquece su maestro")
        res.clases["no es un maestro"] = len(df)
        return res

    def v(fila, rol):
        c = roles.get(rol)
        return "" if c is None or es_vacio(fila[c]) else str(fila[c]).strip()

    def proponer(n, fila, rol, valor, confianza, fuente, porque):
        col = roles.get(rol)
        if col is None:
            res.sin_columna[rol] += 1
            return
        if not es_vacio(fila[col]) and norm(fila[col]) == norm(valor):
            return
        if any(p.fila == n and p.columna == col for p in res.propuestas):
            return
        ident = v(fila, "nombre") or v(fila, "nif")
        res.propuestas.append(Propuesta(n, ident, col, "" if es_vacio(fila[col]) else str(fila[col]).strip(),
                                        valor, confianza, fuente, porque))

    # Las webs se leen antes, en paralelo: lo que tarda es esperar a las que no
    # contestan, y eso no hace falta hacerlo de una en una.
    repetidos_cif, repetidos_tel = set(), set()
    if op.red:
        webs = sorted({w for _, f in df.iterrows() if (w := web_a_leer(f, roles, op))})
        if webs:
            print(f"    leyendo {len(webs)} webs…", flush=True)
            with ThreadPoolExecutor(max_workers=8) as ex:
                leidas = list(ex.map(lambda w: leer_web(red, w), webs))
            # Un CIF o un teléfono que sale en las webs de dos empresas distintas
            # del fichero es de la agencia que las hizo, del alojamiento o del
            # grupo: no es de ninguna de las dos.
            dominios_de = defaultdict(set)
            for d in leidas:
                dom = urllib.parse.urlsplit(d.get("base", "")).netloc.removeprefix("www.")
                for c in d.get("cifs", {}):
                    dominios_de[("cif", c)].add(dom)
                for t in d.get("telefonos", {}):
                    dominios_de[("tel", t)].add(dom)
            repetidos_cif = {c for (k, c), doms in dominios_de.items() if k == "cif" and len(doms) >= 2}
            repetidos_tel = {t for (k, t), doms in dominios_de.items() if k == "tel" and len(doms) >= 2}
            if repetidos_cif or repetidos_tel:
                res.notas.append(f"{len(repetidos_cif)} CIF y {len(repetidos_tel)} teléfonos salen en webs de "
                                 f"varias empresas del fichero (agencia, alojamiento o grupo): no se proponen")

    total = len(df)
    for k, (idx, fila) in enumerate(df.iterrows(), start=1):
        if total > 100 and k % 50 == 0:
            print(f"    … {k}/{total}", flush=True)
        n = int(str(fila[fila_origen]).strip() or 0) if fila_origen and str(fila[fila_origen]).strip().isdigit() \
            else int(idx) + 2
        nif, nombre = v(fila, "nif"), v(fila, "nombre")
        if not nif and not nombre:
            res.clases["vacía"] += 1
            continue
        clase = clasificar(nif, nombre)
        res.clases[clase] += 1
        if clase == "persona" or (clase == "dudoso" and not op.incluir_dudosos):
            continue

        cp, prov, pobl, dire = v(fila, "cp"), v(fila, "provincia"), v(fila, "poblacion"), v(fila, "direccion")

        # 1 · provincia desde el código postal: no sale del ordenador
        cod = provincia_de_cp(cp) if cp else None
        if cod and "provincia" in roles:
            if not prov:
                proponer(n, fila, "provincia", PROVINCIAS[cod], "alta", "código postal",
                         f"Las dos primeras cifras del CP ({cp[:2]}) son la provincia")
            elif codigo_provincia(prov) and codigo_provincia(prov) != cod:
                proponer(n, fila, "provincia", PROVINCIAS[cod], "revisar", "código postal",
                         f"La provincia «{prov}» no cuadra con el CP {cp}: uno de los dos está mal")

        # 2 · CP, población y provincia desde la dirección, con CartoCiudad
        if dire and (not cp or not pobl or not prov) and op.red:
            cand, conf = cartociudad(red, dire, cp, pobl, prov)
            if cand:
                fuente = f"CartoCiudad (IGN): {cand.get('address')}"
                cps = str(cand.get("postalCode") or "").split()
                if not cp and len(cps) == 1:
                    proponer(n, fila, "cp", cps[0], conf, fuente, "Código postal de esa dirección")
                if not pobl and cand.get("muni"):
                    proponer(n, fila, "poblacion", cand["muni"], conf, fuente, "Municipio de esa dirección")
                if not prov and cand.get("province"):
                    proponer(n, fila, "provincia", cand["province"], conf, fuente, "Provincia de esa dirección")
            elif conf.startswith("ningún"):
                res.notas.append(f"Fila {n}: dirección sin candidato fiable ({conf})")

        # 3 · la web de la empresa: la declarada, o la del dominio de su correo
        web = v(fila, "web")
        email = v(fila, "email")
        dominio = email.split("@")[-1].lower().strip(" ;,") if "@" in email else ""
        web_deducida = False
        if not web and dominio and dominio not in CORREO_GENERICO and "." in dominio:
            web, web_deducida = "https://" + dominio, True
        necesita = not nif or not v(fila, "telefono") or not email or web_deducida
        if web and necesita and op.red:
            datos = leer_web(red, web)
            if datos.get("texto"):
                # ¿Es su web? Un correo puede ser del dominio de la gestoría, del
                # grupo o de quien le lleva la informática. Sin su nombre en la
                # web, nada de lo que salga de ella pasa de confianza baja.
                es_suya = nombre_aparece(nombre, datos["texto"])
                aviso = "" if es_suya else "; su nombre no aparece en esa web: puede ser de otra empresa"
                conf_web = "media" if es_suya else "baja"
                if web_deducida:
                    proponer(n, fila, "web", datos["base"], conf_web, datos["base"],
                             f"Dominio de su correo ({dominio}), y la web responde{aviso}")

                # Lo que la web declara de sí misma en schema.org, si es de ella
                org = next((o for o in datos.get("jsonld", []) if o.get("nombre")
                            and nombre_aparece(nombre, o["nombre"])), None)

                if not nif:
                    cif, conf, porque, fuente_cif = elegir_cif(datos, org, nombre, repetidos_cif, es_suya, aviso)
                    if cif:
                        v_ = vies_compara(red, cif, nombre, cp)
                        if v_ and v_["valido"] and v_["nombre"] == "VALID":
                            conf = "alta"
                            porque += " · VIES confirma que el nombre" + (" y el CP" if v_["cp"] == "VALID"
                                                                          else "") + " son de ese CIF"
                        elif v_ and v_["valido"] and v_["nombre"] == "INVALID":
                            conf = "baja"
                            porque += " · VIES dice que ese CIF es de un operador con otro nombre"
                        elif v_ and not v_["valido"]:
                            porque += " · no consta en VIES (no es operador intracomunitario, o no existe)"
                        proponer(n, fila, "nif", cif, conf, fuente_cif, porque)
                    elif porque:
                        res.notas.append(f"Fila {n}: {porque}")

                tels = [t for t, _ in datos["telefonos"].most_common() if t not in repetidos_tel]
                tel_org = re.sub(r"\D", "", (org or {}).get("telefono", ""))[-9:]
                if not v(fila, "telefono") and len(tel_org) == 9 and tel_org[0] in "6789":
                    proponer(n, fila, "telefono", tel_org, conf_web, org["url"],
                             f"El teléfono que declara su web en schema.org{aviso}")
                elif not v(fila, "telefono") and tels:
                    tel = tels[0]
                    proponer(n, fila, "telefono", tel, conf_web, datos["fuente"][("tel", tel)],
                             f"El teléfono que más aparece en su web ({datos['telefonos'][tel]} veces){aviso}")
                if not cp and org and re.fullmatch(r"\d{5}", org.get("cp", "")):
                    proponer(n, fila, "cp", org["cp"], conf_web, org["url"],
                             f"El código postal que declara su web en schema.org{aviso}")
                if not email and datos["emails"]:
                    propios = [e for e in datos["emails"] if e.endswith("@" + urllib.parse.urlsplit(
                        datos["base"]).netloc.removeprefix("www."))] or list(datos["emails"])
                    e = max(propios, key=lambda x: datos["emails"][x])
                    proponer(n, fila, "email", e, conf_web, datos["fuente"][("email", e)],
                             f"Correo publicado en su web{aviso}")

        # 3b · sin CIF en su web: la BDNS, si se pide
        if not nif and op.bdns and op.red and not any(p.fila == n and p.columna == roles.get("nif")
                                                      for p in res.propuestas):
            cif, nombre_bdns, sim = bdns_candidatos(red, nombre)
            if cif:
                conf, porque = "media", f"Beneficiario de una ayuda pública como «{nombre_bdns}»"
                v_ = vies_compara(red, cif, nombre, cp)
                if v_ and v_["valido"] and v_["nombre"] == "VALID":
                    conf, porque = "alta", porque + " · VIES confirma que el nombre es de ese CIF"
                elif v_ and v_["valido"] and v_["nombre"] == "INVALID":
                    conf, porque = "baja", porque + " · VIES dice que ese CIF es de otro nombre"
                proponer(n, fila, "nif", cif, conf, "BDNS (infosubvenciones.es)", porque)
            elif nombre_bdns and "parecidos" in nombre_bdns:
                res.notas.append(f"Fila {n}: {nombre_bdns}")

        # 4 · coherencia de lo que ya hay, sin salir del ordenador
        ident = nombre or nif
        cif_norm = normalizar_id(nif).removeprefix("ES")
        forma = forma_del_nombre(nombre) if nombre else None
        if forma and cif_valido(cif_norm) and cif_norm[0] != LETRA_DE_FORMA[forma]:
            res.comprobaciones.append((n, ident, "Letra del CIF y forma jurídica",
                                       f"REVISAR: el nombre dice {forma} y un CIF que empieza por {cif_norm[0]} "
                                       f"no es de {forma} (le tocaría {LETRA_DE_FORMA[forma]})",
                                       "Orden EHA/451/2008 (letras del NIF de entidades)"))
        provs = provincias_de_telefono(v(fila, "telefono"))
        if cod and provs and cod not in provs:
            res.comprobaciones.append((n, ident, "Prefijo del teléfono y CP",
                                       f"AVISO: el fijo es de {', '.join(PROVINCIAS[p] for p in sorted(provs))} y el "
                                       f"CP de {PROVINCIAS[cod]}. Puede ser la sede; o uno de los dos está mal",
                                       "Plan nacional de numeración"))

        # 5 · VIES para los CIF que ya tiene, si se pide: ¿es suyo ese CIF?
        if op.vies and nif and cif_valido(cif_norm) and op.red:
            v_ = vies_compara(red, cif_norm, nombre, cp) if nombre else None
            res.vies.append((n, nif, None if v_ is None else v_["valido"]))
            if v_ and v_["valido"] and v_["nombre"] == "INVALID":
                # Señal débil: VIES solo casa la razón social casi completa. En un
                # maestro real, 111 de 390 no casaban, y eran sobre todo nombres
                # comerciales («Hotel…», «… Suites») y razones sociales recortadas.
                res.comprobaciones.append((n, ident, "VIES: nombre del CIF",
                                           "AVISO: VIES no reconoce este nombre para ese CIF. Casi siempre es un "
                                           "nombre comercial o una razón social recortada; si el nombre es la "
                                           "razón social completa, el CIF puede ser de otra empresa", "VIES"))
            elif v_ and v_["valido"] and v_["cp"] == "INVALID":
                res.comprobaciones.append((n, ident, "VIES: CP del CIF",
                                           "AVISO: el nombre cuadra pero el CP no es el que tiene VIES: "
                                           "¿domicilio fiscal distinto?", "VIES"))
    return res


def elegir_cif(datos: dict, org: dict | None, nombre: str, repetidos: set, es_suya: bool, aviso: str) -> tuple:
    """(cif, confianza, por qué, fuente) o (None, None, motivo, None). Primero
    lo que la web declara de sí misma en schema.org; después el del bloque que
    identifica al titular. Nunca el de la agencia, ni uno que sale en las webs
    de varias empresas del fichero."""
    if org:
        c = normalizar_id(org.get("cif", "")).removeprefix("ES")
        if cif_valido(c) and c not in repetidos:
            return c, "alta", "El CIF que declara su web en schema.org, con su nombre", org["url"]
    candidatos = [c for c in datos["cifs"] if c not in repetidos]
    if not candidatos:
        agencias = [c for c, s in datos.get("senales", {}).items() if s["agencia"]]
        if agencias or any(c in repetidos for c in datos["cifs"]):
            return None, None, "el único CIF de su web es de la agencia o sale en otras webs: no se propone", None
        return None, None, "", None
    if len(candidatos) > 1:
        puntos = {c: datos["senales"][c]["puntos"] for c in candidatos}
        mejor = max(puntos.values())
        del_titular = [c for c in candidatos if puntos[c] == mejor]
        if mejor == 0 or len(del_titular) != 1:
            return None, None, f"su web lleva {len(candidatos)} CIF distintos; no se propone ninguno", None
        candidatos = del_titular
    cif = candidatos[0]
    s = datos["senales"][cif]
    fuente = datos["fuente"][("cif", cif)]
    if nombre_aparece(nombre, datos["contexto"][cif]):
        return cif, "alta", "El CIF de su web, junto a su nombre", fuente
    if s["registro"] or s["titular"]:
        return cif, "media" if es_suya else "baja", \
            "El CIF que su web da como titular (aviso legal, datos registrales), sin su nombre al lado" + aviso, fuente
    return cif, "media" if es_suya else "baja", "El único CIF de su web, pero no junto a su nombre" + aviso, fuente


# ---------------------------------------------------------------------------
# Escribir
# ---------------------------------------------------------------------------

COLUMNAS = ["Fila", "Identificador", "Columna", "Valor actual", "💡 Valor propuesto", "Confianza",
            "Fuente", "Por qué", "✏️ Valor corregido", "Decisión", "Notas"]
ANCHOS = [8, 30, 16, 22, 28, 11, 40, 44, 22, 22, 26]


def nombre_de_pestana(entidad: str, usados: set) -> str:
    """Como las nombra auditar-datos.py, para que limpiar-datos.py las reconozca."""
    base = re.sub(r"[\\/*?:\[\]]", "-", entidad)[:28] or "Entidad"
    nombre, n = base, 2
    while nombre in usados:
        nombre = f"{base[:26]}-{n}"
        n += 1
    usados.add(nombre)
    return nombre


def comprobar_destino(ruta: Path, permitir_git: bool) -> None:
    if CARPETA_ORIGEN in {norm(p) for p in ruta.resolve().parts}:
        morir(f"No escribo dentro de 02-Datos-Origen/: el origen no se toca. Salida: {ruta}")
    if permitir_git:
        return
    carpeta = ruta.parent if ruta.parent.exists() else Path.cwd()
    try:
        dentro = subprocess.run(["git", "-C", str(carpeta), "rev-parse", "--is-inside-work-tree"],
                                capture_output=True, text=True)
        if dentro.returncode != 0 or dentro.stdout.strip() != "true":
            return
        ignorado = subprocess.run(["git", "-C", str(carpeta), "check-ignore", "-q", str(ruta.resolve())],
                                  capture_output=True)
    except FileNotFoundError:
        return
    if ignorado.returncode != 0:
        morir(f"La salida está dentro de un repositorio git y no está ignorada: {ruta}\n"
              f"  Los datos del cliente no se publican. Escribe fuera, o añádela al .gitignore.")


def escribir(resultados: list, ruta: Path, red: Red, op) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    azul, blanco = PatternFill("solid", fgColor="1F3864"), Font(bold=True, color="FFFFFF")
    verde = PatternFill("solid", fgColor="E8F5E9")
    tonos = {"alta": "E8F5E9", "media": "FFF3CD", "baja": "FCE4E4", "revisar": "FCE4E4"}

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Datos que faltan · propuestas"
    ws["A1"].font = Font(bold=True, size=16, color="1F3864")
    ws["A2"] = f"Generado el {datetime.now():%d/%m/%Y} con enriquecer-datos.py {VERSION}"
    ws["A2"].font = Font(italic=True, color="666666")
    ws["A4"] = ("Cada fila es un dato que falta y que hemos encontrado en una fuente pública. No se ha "
                "cambiado nada: si la propuesta es buena, marca «Se acepta la propuesta». Si no, escribe el "
                "valor bueno en «Valor corregido» o márcala como «Está bien así».")
    ws["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A4:F4")
    ws.row_dimensions[4].height = 45
    f = 6
    for i, cab in enumerate(["Entidad", "Filas", "Empresas", "Personas (no se buscan)", "Dudosas",
                             "Propuestas", "Comprobaciones"], start=1):
        c = ws.cell(f, i, cab)
        c.fill, c.font = azul, blanco
    for r in resultados:
        f += 1
        for i, val in enumerate([r.entidad, r.filas, r.clases["empresa"], r.clases["persona"],
                                 r.clases["dudoso"], len(r.propuestas), len(r.comprobaciones)], start=1):
            ws.cell(f, i, val)
    f += 2
    ws.cell(f, 1, "Consultas hechas").font = Font(bold=True)
    for tipo, n in red.consultas.items():
        f += 1
        ws.cell(f, 1, tipo)
        ws.cell(f, 2, n)
    for texto, conjunto in (("Webs con protección contra robots (no se leen)", red.antibot),
                            ("Webs que pidieron parar (429/503)", red.abandonados)):
        if conjunto:
            f += 1
            ws.cell(f, 1, texto)
            ws.cell(f, 2, len(conjunto))
    notas = [(r.entidad, x) for r in resultados for x in r.notas] + \
            [(r.entidad, f"{n} datos de «{rol}» encontrados, pero el fichero no tiene esa columna")
             for r in resultados for rol, n in r.sin_columna.items()]
    if notas:
        f += 2
        ws.cell(f, 1, "Notas para Tazuke").font = Font(bold=True, color="9C6500")
        for ent, x in notas:
            f += 1
            ws.cell(f, 1, ent)
            ws.cell(f, 2, x)
    for i, a in enumerate([34, 12, 12, 22, 12, 12], start=1):
        ws.column_dimensions[get_column_letter(i)].width = a

    usados = {"Resumen", "VIES", "Comprobaciones"}
    for r in resultados:
        if not r.propuestas:
            continue
        h = wb.create_sheet(nombre_de_pestana(r.entidad, usados))
        h.sheet_view.showGridLines = False
        for i, cab in enumerate(COLUMNAS, start=1):
            c = h.cell(1, i, cab)
            c.fill, c.font = azul, blanco
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        orden = {"revisar": 0, "alta": 1, "media": 2, "baja": 3}
        for k, p in enumerate(sorted(r.propuestas, key=lambda x: (orden.get(x.confianza, 9), x.fila)), start=2):
            for i, val in enumerate([p.fila, p.identificador, p.columna, p.valor_actual, p.propuesto,
                                     p.confianza.upper(), p.fuente, p.porque, "", "", ""], start=1):
                c = h.cell(k, i, val)
                c.alignment = Alignment(vertical="top", wrap_text=i in (7, 8))
            h.cell(k, 6).fill = PatternFill("solid", fgColor=tonos.get(p.confianza, "FFFFFF"))
            h.cell(k, 9).fill = verde
        dv = DataValidation(type="list", allow_blank=True,
                            formula1='"Se acepta la propuesta,Corregido,Está bien así,Lo decide Tazuke"')
        h.add_data_validation(dv)
        dv.add(f"J2:J{len(r.propuestas) + 1}")
        h.freeze_panes = "A2"
        h.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}{len(r.propuestas) + 1}"
        for i, a in enumerate(ANCHOS, start=1):
            h.column_dimensions[get_column_letter(i)].width = a

    if any(r.comprobaciones for r in resultados):
        # Lo que no es una propuesta sino una duda sobre lo que ya hay: no se
        # aplica nada desde aquí. Es para la siguiente conversación con el cliente.
        h = wb.create_sheet("Comprobaciones")
        h.append(["Entidad", "Fila", "Identificador", "Qué se ha comprobado", "Resultado", "Fuente"])
        for i in range(1, 7):
            h.cell(1, i).fill, h.cell(1, i).font = azul, blanco
        for r in resultados:
            for fila in sorted(r.comprobaciones, key=lambda x: (0 if x[3].startswith("REVISAR") else 1, x[0])):
                h.append([r.entidad, *fila])
        for i, a in enumerate([24, 8, 30, 28, 80, 30], start=1):
            h.column_dimensions[get_column_letter(i)].width = a
        h.freeze_panes = "A2"
    if any(r.vies for r in resultados):
        h = wb.create_sheet("VIES")
        h.append(["Entidad", "Fila", "CIF", "¿Operador intracomunitario?"])
        for r in resultados:
            for n, cif, ok in r.vies:
                h.append([r.entidad, n, cif, "sí" if ok else ("no" if ok is False else "sin respuesta")])
    wb.save(ruta)


# ---------------------------------------------------------------------------
# Orquestación
# ---------------------------------------------------------------------------

@dataclass
class Opciones:
    red: bool = True
    vies: bool = False
    bdns: bool = False
    incluir_dudosos: bool = False


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(prog="enriquecer-datos.py",
                                 description="Propone los datos que le faltan a un maestro de empresas.")
    ap.add_argument("ficheros", nargs="*", type=Path)
    ap.add_argument("--salida", type=Path)
    ap.add_argument("--columna", action="append", metavar="ROL=COLUMNA")
    ap.add_argument("--sin-red", action="store_true")
    ap.add_argument("--vies", action="store_true")
    ap.add_argument("--bdns", action="store_true")
    ap.add_argument("--incluir-dudosos", action="store_true")
    ap.add_argument("--simular", action="store_true")
    ap.add_argument("--permitir-git", action="store_true")
    ap.add_argument("--version", action="version", version=f"enriquecer-datos.py {VERSION}")
    a = ap.parse_args(argv)

    try:
        import openpyxl  # noqa: F401
        import pandas  # noqa: F401
    except ImportError:
        morir("Faltan dependencias:  pip install pandas openpyxl")

    forzados = {}
    for par in a.columna or []:
        if "=" not in par:
            morir(f"--columna espera ROL=COLUMNA, y llegó «{par}»")
        rol, col = par.split("=", 1)
        if norm(rol) not in ROLES:
            morir(f"Rol desconocido «{rol}». Los que hay: {', '.join(ROLES)}")
        forzados[norm(rol)] = col.strip()

    if a.ficheros:
        ficheros = a.ficheros
        salida = a.salida or Path("propuestas")
    else:
        if not ENTRADA_METODO.exists():
            morir(f"Sin ficheros y sin {ENTRADA_METODO}/. Pasa primero limpiar-datos.py, o indica ficheros.")
        ficheros = sorted(p for p in ENTRADA_METODO.glob("*.xlsx")
                          if not p.name.startswith(("Bitacora-", "~$")))
        salida = a.salida or SALIDA_METODO
        if not ficheros:
            morir(f"{ENTRADA_METODO}/ está vacía. Pasa primero limpiar-datos.py.")
    for f in ficheros:
        if not f.exists():
            morir(f"No existe: {f}")

    op = Opciones(red=not a.sin_red, vies=a.vies, bdns=a.bdns, incluir_dudosos=a.incluir_dudosos)
    ruta = salida / f"Enriquecimiento-{datetime.now():%Y-%m-%d}.xlsx"
    cache_path = None if a.simular else salida / ".cache-enriquecer.json"
    if not a.simular:
        salida.mkdir(parents=True, exist_ok=True)
        comprobar_destino(ruta, a.permitir_git)
        comprobar_destino(cache_path, a.permitir_git)
    red = Red(cache_path, activa=op.red)

    print(f"enriquecer-datos.py {VERSION} · {len(ficheros)} ficheros · red {'sí' if op.red else 'no'}")
    resultados = []
    try:
        for f in ficheros:
            hojas = cargar_hojas(f)
            for hoja, df in hojas.items():
                if not len(df.columns):
                    continue
                entidad = hoja if not a.ficheros else (f.stem if len(hojas) == 1 else f"{f.stem} · {hoja}")
                roles = asignar_roles([str(c) for c in df.columns], forzados)
                r = enriquecer(df.astype(str), entidad, str(f), roles, red, op)
                resultados.append(r)
                print(f"  {entidad}: {r.filas} filas · {r.clases['empresa']} empresas · "
                      f"{r.clases['persona']} personas (no se buscan) · {len(r.propuestas)} propuestas")
    finally:
        if not a.simular:
            red.guardar()

    if red.consultas:
        print("  Consultas: " + ", ".join(f"{t} {n}" for t, n in red.consultas.items()))
    if a.simular:
        print("  (simulación: no se ha escrito nada)")
        return 0
    escribir(resultados, ruta, red, op)
    print(f"  Escrito: {ruta}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
