# -*- coding: utf-8 -*-
"""Pruebas de enriquecer-datos.py. Sin red: las respuestas de CartoCiudad y de
las webs son inventadas, igual que todos los datos. En este repositorio no
entra nunca un dato real de un cliente."""

import importlib.util
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
spec = importlib.util.spec_from_file_location("enriquecer_datos", RAIZ / "enriquecer-datos.py")
ed = importlib.util.module_from_spec(spec)
sys.modules["enriquecer_datos"] = ed
spec.loader.exec_module(ed)

CARTO = [{
    "type": "portal", "address": "CALLE MAYOR 12, Villaficticia", "tip_via": "CALLE",
    "portalNumber": 12, "postalCode": "29100", "muni": "Villaficticia", "poblacion": "Villaficticia",
    "province": "Málaga", "provinceCode": "29",
}, {
    "type": "portal", "address": "CALLE MAYORAL 2, Otro Sitio", "tip_via": "CALLE",
    "portalNumber": 2, "postalCode": "43892", "muni": "Otro Sitio", "poblacion": "Otro Sitio",
    "province": "Tarragona", "provinceCode": "43",
}]
PORTADA = '<html><a href="/aviso-legal">Aviso legal</a><p>Llámanos: 952 000 111</p></html>'
AVISO = ('<html><p>TALLERES FICTICIOS DEL SUR, S.L. · CIF B-12.345.674 · '
         'Calle Mayor 12, Villaficticia · info@talleresficticios.example</p></html>')


class RedFalsa(ed.Red):
    """Contesta de memoria, y apunta lo que le preguntan."""

    def __init__(self, vies=None):
        super().__init__(None, activa=True)
        self.preguntas = []
        self.respuesta_vies = vies      # lo que contestaría VIES, o None: sin respuesta

    def post_json(self, url, datos, tipo):
        self.preguntas.append(url)
        return self.respuesta_vies

    def descargar(self, url, tipo, max_bytes=600_000):
        self.preguntas.append(url)
        self.consultas[tipo] += 1
        if "cartociudad" in url:
            return "callback(" + json.dumps(CARTO) + ")"
        if url.endswith("/robots.txt"):
            return "User-agent: *\nDisallow: /privado\n"
        if url.rstrip("/").endswith("talleresficticios.example"):
            return PORTADA
        if "aviso-legal" in url:
            return AVISO
        if url.rstrip("/").endswith("gestoria.example"):
            return '<p>GESTORÍA INVENTADA SL · CIF A58818501 · 952 999 888</p>'
        return None


class Piezas(unittest.TestCase):
    def test_clasificar(self):
        self.assertEqual(ed.clasificar("", "Talleres Norte S.L."), "empresa")
        self.assertEqual(ed.clasificar("B12345674", ""), "empresa")
        self.assertEqual(ed.clasificar("12345678Z", "Ana"), "persona")
        self.assertEqual(ed.clasificar("X1234567L", ""), "persona")
        self.assertEqual(ed.clasificar("", "Ana Pérez"), "dudoso")

    def test_provincias(self):
        self.assertEqual(ed.provincia_de_cp("4001"), "04")
        self.assertEqual(ed.provincia_de_cp("99000"), None)
        self.assertEqual(ed.codigo_provincia("VIZCAYA"), "48")
        self.assertEqual(ed.codigo_provincia("Málaga"), "29")

    def test_cif_y_telefono_en_texto(self):
        t = "CIF: B-12.345.674, tel +34 600-111-222 y 952 12 34 56".upper()
        self.assertEqual(["".join(m.groups()) for m in ed.RE_CIF_TEXTO.finditer(t)], ["B12345674"])
        self.assertTrue(ed.cif_valido("B12345674"))
        self.assertFalse(ed.cif_valido("B12345675"))
        tels = [ed.re.sub(r"\D", "", m.group(0))[-9:] for m in ed.RE_TEL.finditer(t)]
        self.assertEqual(tels, ["600111222", "952123456"])

    def test_calle(self):
        self.assertEqual(ed.nucleo_calle("C/ Marqués de Larios, 5 - 2ºB"), "marques larios")
        self.assertEqual(ed.numero_portal("Avda. Andalucía 12, 1º"), 12)


class Cartociudad(unittest.TestCase):
    def test_solo_vale_si_cuadra(self):
        red = RedFalsa()
        cand, conf = ed.cartociudad(red, "C/ Mayor 12", "", "Villaficticia", "")
        self.assertEqual(cand["postalCode"], "29100")
        self.assertEqual(conf, "alta")
        # La misma calle con un CP de otra provincia: ningún candidato vale
        cand, motivo = ed.cartociudad(red, "C/ Mayor 12", "08001", "", "")
        self.assertIsNone(cand)
        self.assertIn("ningún candidato", motivo)


def escribir_xlsx(ruta, filas):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Proveedores"
    for f in filas:
        ws.append(f)
    wb.save(ruta)


class DePuntaAPunta(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)
        self.fichero = self.d / "proveedores.xlsx"
        escribir_xlsx(self.fichero, [
            ["Nombre", "CIF", "Dirección", "CP", "Población", "Provincia", "Teléfono", "Email", "Web"],
            ["Talleres Ficticios del Sur, S.L.", "", "C/ Mayor 12", "", "Villaficticia", "", "",
             "", "www.talleresficticios.example"],
            ["Ana Pérez Inventada", "12345678Z", "C/ Mayor 12", "", "Villaficticia", "", "", "", ""],
            ["Otra Empresa Inventada SA", "A00000000", "", "29001", "", "Barcelona", "600000000",
             "hola@otra.example", ""],
        ])
        self.cwd = os.getcwd()
        os.chdir(self.d)

    def tearDown(self):
        os.chdir(self.cwd)
        self.tmp.cleanup()

    def test_propuestas(self):
        import pandas as pd
        df = pd.read_excel(self.fichero, dtype=str, keep_default_na=False)
        roles = ed.asignar_roles(list(df.columns), {})
        red = RedFalsa()
        r = ed.enriquecer(df, "Proveedores", str(self.fichero), roles, red, ed.Opciones())
        prop = {(p.fila, p.columna): p for p in r.propuestas}

        self.assertEqual(prop[(2, "CIF")].propuesto, "B12345674")
        self.assertEqual(prop[(2, "CIF")].confianza, "alta", "el nombre está junto al CIF")
        self.assertEqual(prop[(2, "CP")].propuesto, "29100")
        self.assertEqual(prop[(2, "Provincia")].propuesto, "Málaga")
        self.assertEqual(prop[(2, "Teléfono")].propuesto, "952000111")
        self.assertEqual(prop[(4, "Provincia")].confianza, "revisar", "Barcelona no cuadra con el CP 29001")

        self.assertFalse(any(p.fila == 3 for p in r.propuestas), "una persona física no se busca")
        self.assertEqual(r.clases["persona"], 1)
        self.assertFalse(any("privado" in u for u in red.preguntas))

    def test_excel_con_formato_de_auditoria(self):
        codigo = ed.main([str(self.fichero), "--salida", str(self.d / "p"), "--sin-red", "--permitir-git"])
        self.assertEqual(codigo, 0)
        from openpyxl import load_workbook
        wb = load_workbook(next((self.d / "p").glob("Enriquecimiento-*.xlsx")))
        ws = wb["proveedores"]
        cab = [c.value for c in ws[1]]
        for imprescindible in ("Fila", "Columna", "Valor actual", "💡 Valor propuesto", "✏️ Valor corregido",
                               "Decisión"):
            self.assertIn(imprescindible, cab)
        self.assertTrue(any(r[4] == "Barcelona" or r[4] == "Málaga"
                            for r in ws.iter_rows(min_row=2, values_only=True)))

    def test_web_de_otra_empresa(self):
        """El correo es de la gestoría: su CIF no puede salir con confianza alta."""
        import pandas as pd
        df = pd.DataFrame([{"Nombre": "Bodegas Inventadas SL", "CIF": "", "Email": "bodegas@gestoria.example",
                            "Teléfono": ""}])
        r = ed.enriquecer(df, "P", "p.xlsx", ed.asignar_roles(list(df.columns), {}), RedFalsa(), ed.Opciones())
        prop = {p.columna: p for p in r.propuestas}
        self.assertEqual(prop["CIF"].propuesto, "A58818501")
        self.assertEqual(prop["CIF"].confianza, "baja")
        self.assertIn("otra empresa", prop["CIF"].porque)
        self.assertEqual(prop["Teléfono"].confianza, "baja")

    def test_aviso_legal_al_pie(self):
        """Con muchos enlaces delante, el aviso legal se lee igual, y los
        ficheros estáticos no gastan visitas."""
        enlaces = "".join(f'<a href="/empresa/{i}">x</a>' for i in range(500))
        enlaces += '<a href="/wp-content/plugins/legal.css">c</a><a href="/aviso-legal">Aviso legal</a>'

        class Red2(RedFalsa):
            def descargar(self, url, tipo, max_bytes=600_000):
                self.preguntas.append(url)
                if url.endswith("robots.txt"):
                    return ""
                if url.endswith("aviso-legal"):
                    return AVISO
                return "<html>" + enlaces + "</html>"
        red = Red2()
        d = ed.leer_web(red, "https://talleresficticios.example")
        self.assertIn("B12345674", d["cifs"])
        self.assertFalse(any(u.endswith(".css") for u in red.preguntas))

    def test_nombre_palabras_enteras(self):
        self.assertFalse(ed.nombre_aparece("Ron SL", "Ronda del sur"))
        self.assertTrue(ed.nombre_aparece("Bodegas Inventadas, S.L.", "BODEGAS INVENTADAS SL, CIF..."))

    def test_sin_red_no_pregunta(self):
        red = ed.Red(None, activa=False)
        self.assertIsNone(red.get("https://example.invalid/", "web"))
        self.assertEqual(sum(red.consultas.values()), 0)


class RedWebs(RedFalsa):
    """Varias webs inventadas, cada una con su trampa."""
    PAGINAS = {
        "agencia.example": '<p>Hostelería Inventada SL</p><footer><p>Diseño web por Agencia Falsa SL · '
                           'CIF B12345674</p></footer>',
        "dos-cifs.example": '<p>Cliente nuestro: Otra SA, CIF A58818501</p>'
                            '<p>Titular: Dos Cifs Inventada SL · CIF B87654323 · Registro Mercantil de '
                            'Málaga, Tomo 1, Folio 2, Hoja MA-3</p>',
        "jsonld.example": '<script type="application/ld+json">{"@context":"https://schema.org",'
                          '"@type":"LocalBusiness","name":"Jsonld Inventada SL","telephone":"+34 952 111 222",'
                          '"address":{"@type":"PostalAddress","postalCode":"29001"}}</script><p>Hola</p>',
        "a.example": '<p>Uno Inventada SL</p><p>Teléfono 952 333 444</p>',
        "b.example": '<p>Dos Inventada SL</p><p>Teléfono 952 333 444</p>',
    }

    def descargar(self, url, tipo, max_bytes=600_000):
        self.preguntas.append(url)
        if url.endswith("robots.txt"):
            return ""
        host = ed.urllib.parse.urlsplit(url).netloc
        return self.PAGINAS.get(host)


class Aprendizajes(unittest.TestCase):
    """Lo que salió de revisar herramientas ajenas el 2026-09-27."""

    def enriquecer(self, filas, red):
        import pandas as pd
        df = pd.DataFrame(filas)
        return ed.enriquecer(df, "P", "p.xlsx", ed.asignar_roles(list(df.columns), {}), red, ed.Opciones())

    def test_cif_de_la_agencia_no_se_propone(self):
        r = self.enriquecer([{"Nombre": "Hostelería Inventada SL", "CIF": "", "Web": "agencia.example",
                              "Teléfono": "600000000", "Email": "x@y.example"}], RedWebs())
        self.assertFalse(any(p.columna == "CIF" for p in r.propuestas))

    def test_de_dos_cifs_el_del_titular(self):
        r = self.enriquecer([{"Nombre": "Otra Empresa Inventada SL", "CIF": "", "Web": "dos-cifs.example",
                              "Teléfono": "600000000", "Email": "x@y.example"}], RedWebs())
        prop = {p.columna: p for p in r.propuestas}
        self.assertEqual(prop["CIF"].propuesto, "B87654323")
        self.assertIn("titular", prop["CIF"].porque)

    def test_jsonld(self):
        r = self.enriquecer([{"Nombre": "Jsonld Inventada SL", "CIF": "B12345674", "Web": "jsonld.example",
                              "Teléfono": "", "Email": "", "CP": ""}], RedWebs())
        prop = {p.columna: p for p in r.propuestas}
        self.assertEqual(prop["Teléfono"].propuesto, "952111222")
        self.assertIn("schema.org", prop["Teléfono"].porque)
        self.assertEqual(prop["CP"].propuesto, "29001")

    def test_telefono_repetido_en_dos_webs(self):
        r = self.enriquecer([{"Nombre": "Uno Inventada SL", "CIF": "B12345674", "Web": "a.example", "Teléfono": ""},
                             {"Nombre": "Dos Inventada SL", "CIF": "B87654323", "Web": "b.example", "Teléfono": ""}],
                            RedWebs())
        self.assertFalse(any(p.columna == "Teléfono" for p in r.propuestas))

    def test_vies_sube_y_baja_la_confianza(self):
        fila = [{"Nombre": "Talleres Ficticios del Sur SL", "CIF": "", "Web": "", "Email": "",
                 "Teléfono": "", "Dirección": "", "CP": "29100"}]
        fila[0]["Web"] = "www.talleresficticios.example"
        ok = self.enriquecer(fila, RedFalsa(vies={"valid": True, "traderNameMatch": "VALID",
                                                  "traderPostalCodeMatch": "VALID"}))
        cif = next(p for p in ok.propuestas if p.columna == "CIF")
        self.assertEqual(cif.confianza, "alta")
        self.assertIn("VIES confirma", cif.porque)
        mal = self.enriquecer(fila, RedFalsa(vies={"valid": True, "traderNameMatch": "INVALID"}))
        cif = next(p for p in mal.propuestas if p.columna == "CIF")
        self.assertEqual(cif.confianza, "baja")

    def test_comprobaciones(self):
        r = self.enriquecer([{"Nombre": "Grande Inventada, S.A.", "CIF": "B12345674", "CP": "29001",
                              "Teléfono": "912345678", "Provincia": "Málaga"}], RedFalsa())
        que = " ".join(c[2] for c in r.comprobaciones)
        self.assertIn("Letra del CIF", que)
        self.assertIn("Prefijo del teléfono", que)

    def test_nombre_corto_para_vies(self):
        self.assertEqual(ed.nombre_corto("Telefónica, Sociedad Anónima"), "TELEFONICA SA")
        self.assertEqual(ed.nombre_corto("Talleres Pérez S. L."), "TALLERES PEREZ SL")
        self.assertEqual(ed.nombre_corto("Renfe Viajeros S.M.E., S.A."), "RENFE VIAJEROS SME SA")
        self.assertEqual(ed.nombre_corto("Gruas Palli S.A.U."), "GRUAS PALLI SAU")

    def test_texto_por_bloques(self):
        t = ed.texto_de_html("<div>Uno</div><p>Dos &amp; tres</p><br>cuatro")
        self.assertEqual(t.split("\n"), ["Uno", "Dos & tres", "cuatro"])


class RedBdns(RedFalsa):
    TERCEROS = {"terceros": [
        {"id": 1, "descripcion": "B26613679 - BLINDITEX INVENTADA, S.L."},
        {"id": 2, "descripcion": "12345678Z - INDITEX INVENTADA PERSONA"},
        {"id": 3, "descripcion": "A58818501 - INDITEX INVENTADA S.A."},
    ]}

    def descargar(self, url, tipo, max_bytes=600_000):
        self.preguntas.append(url)
        return json.dumps(self.TERCEROS) if "bdnstrans" in url else None


class Bdns(unittest.TestCase):
    def test_consultas(self):
        self.assertEqual(ed.consultas_bdns("Telefónica de España SAU")[0], "TELEFONICA DE ESPA")
        self.assertEqual(ed.consultas_bdns("Mercadona S.A."), ["MERCADONA"])

    def test_solo_el_bueno(self):
        cif, nombre, sim = ed.bdns_candidatos(RedBdns(), "Inditex Inventada SA")
        self.assertEqual(cif, "A58818501")

    def test_ni_trozos_ni_otra_forma(self):
        # «Blinditex» contiene «inditex», pero no es la palabra; y una SL no es una SA
        cif, motivo, _ = ed.bdns_candidatos(RedBdns(), "Blinditex Inventada SA")
        self.assertIsNone(cif)

    def test_propuesta_con_vies(self):
        import pandas as pd
        red = RedBdns(vies={"valid": True, "traderNameMatch": "VALID"})
        df = pd.DataFrame([{"Nombre": "Inditex Inventada SA", "CIF": "", "CP": ""}])
        op = ed.Opciones(bdns=True)
        r = ed.enriquecer(df, "P", "p.xlsx", ed.asignar_roles(list(df.columns), {}), red, op)
        p = next(x for x in r.propuestas if x.columna == "CIF")
        self.assertEqual((p.propuesto, p.confianza), ("A58818501", "alta"))

    def test_sin_bdns_no_pregunta(self):
        import pandas as pd
        red = RedBdns()
        df = pd.DataFrame([{"Nombre": "Inditex Inventada SA", "CIF": ""}])
        ed.enriquecer(df, "P", "p.xlsx", ed.asignar_roles(list(df.columns), {}), red, ed.Opciones())
        self.assertFalse(any("bdnstrans" in u for u in red.preguntas))


if __name__ == "__main__":
    unittest.main()
